"""MCP server: Office documents (docx / xlsx / pptx) plus CSV table helpers."""
import csv
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from mcpserver import Server  # noqa: E402

import openpyxl  # noqa: E402
from docx import Document  # noqa: E402
from docx.table import Table  # noqa: E402
from docx.text.paragraph import Paragraph  # noqa: E402
from pptx import Presentation  # noqa: E402

MAX_ROWS = 200
TEMP = os.environ.get("TEMP") or os.environ.get("TMP") or "."
SAMPLE_DIR = os.path.join(TEMP, "aiyu_mcp_samples")

srv = Server("office")


def _in_path(path):
    """Resolve an input path; a missing input becomes a clear isError result."""
    p = os.path.abspath(os.path.expanduser(str(path)))
    if not os.path.isfile(p):
        raise FileNotFoundError(path)
    return p


def _out_path(path):
    """Resolve an output path, creating its parent directories."""
    p = os.path.abspath(os.path.expanduser(str(path)))
    os.makedirs(os.path.dirname(p) or ".", exist_ok=True)
    return p


def _cell(value):
    return "" if value is None else str(value)


def _block(rows):
    return "\n".join(" | ".join(_cell(v) for v in row) for row in rows)


def _coerce(text):
    """Best-effort int/float typing so CSV numbers stay numeric in Excel."""
    s = str(text).strip()
    try:
        return int(s)
    except ValueError:
        pass
    try:
        return float(s)
    except ValueError:
        return text
@srv.tool("docx_create", "Create a .docx with optional heading, paragraphs and a table.",
          {"type": "object", "properties": {"path": {"type": "string"}, "title": {"type": "string", "default": ""}, "paragraphs": {"type": "array", "items": {"type": "string"}}, "table": {"type": "array", "items": {"type": "array"}}}, "required": ["path"]})
def docx_create(path, title="", paragraphs=[], table=[[]]):
    p = _out_path(path)
    doc = Document()
    if title:
        doc.add_heading(str(title), level=1)
    for text in paragraphs or []:
        doc.add_paragraph(str(text))
    rows = [list(r) for r in (table or []) if r]
    if rows:
        cols = max(len(r) for r in rows)
        tbl = doc.add_table(rows=len(rows), cols=cols)
        tbl.style = "Table Grid"
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                tbl.cell(i, j).text = _cell(val)
    doc.save(p)
    return "created %s (title=%r paragraphs=%d table=%dx%d)" % (p, title, len(paragraphs or []), len(rows), len(rows[0]) if rows else 0)


@srv.tool("docx_read", "Read a .docx as text; headings and tables are marked.",
          {"type": "object", "properties": {"path": {"type": "string"}}, "required": ["path"]})
def docx_read(path):
    p = _in_path(path)
    doc = Document(p)
    out = []
    for child in doc.element.body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            para = Paragraph(child, doc)
            text = para.text.strip()
            if not text:
                continue
            style = para.style.name if para.style is not None else ""
            style = style or ""
            if style.lower().startswith("heading"):
                out.append("[%s] %s" % (style.upper().replace(" ", ""), text))
            elif style.lower() == "title":
                out.append("[TITLE] %s" % text)
            else:
                out.append(text)
        elif tag == "tbl":
            tbl = Table(child, doc)
            out.append("[TABLE %dx%d]" % (len(tbl.rows), len(tbl.columns)))
            for row in tbl.rows:
                out.append("  " + " | ".join(_cell(c.text.strip()) for c in row.cells))
    return "path=%s lines=%d\n%s" % (p, len(out), "\n".join(out) or "(empty document)")


@srv.tool("docx_append", "Append paragraphs to an existing .docx.",
          {"type": "object", "properties": {"path": {"type": "string"}, "paragraphs": {"type": "array", "items": {"type": "string"}}}, "required": ["path"]})
def docx_append(path, paragraphs=[]):
    p = _in_path(path)
    doc = Document(p)
    for text in paragraphs or []:
        doc.add_paragraph(str(text))
    doc.save(p)
    return "appended %d paragraph(s) -> %s" % (len(paragraphs or []), p)
@srv.tool("xlsx_create", "Create an .xlsx workbook from a list of rows.",
          {"type": "object", "properties": {"path": {"type": "string"}, "rows": {"type": "array", "items": {"type": "array"}}, "sheet": {"type": "string", "default": "Sheet1"}}, "required": ["path"]})
def xlsx_create(path, rows=[[]], sheet="Sheet1"):
    p = _out_path(path)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = str(sheet or "Sheet1")
    count = 0
    for row in rows or []:
        ws.append(list(row))
        count += 1
    wb.save(p)
    return "created %s (sheet=%s rows=%d)" % (p, ws.title, count)


@srv.tool("xlsx_read", "Read a sheet as 'a | b | c' lines (default: first sheet).",
          {"type": "object", "properties": {"path": {"type": "string"}, "sheet": {"type": "string", "default": ""}, "max_rows": {"type": "integer", "default": MAX_ROWS}}, "required": ["path"]})
def xlsx_read(path, sheet="", max_rows=MAX_ROWS):
    p = _in_path(path)
    limit = max(1, int(max_rows))
    wb = openpyxl.load_workbook(p, data_only=True, read_only=True)
    try:
        ws = wb[str(sheet)] if sheet and str(sheet) in wb.sheetnames else wb[wb.sheetnames[0]]
        lines = []
        total = 0
        for row in ws.iter_rows(values_only=True):
            total += 1
            if len(lines) < limit:
                vals = list(row)
                while vals and vals[-1] is None:
                    vals.pop()
                lines.append(" | ".join(_cell(v) for v in vals))
        title = ws.title
    finally:
        wb.close()
    if total > limit:
        lines.append("...[truncated at %d of %d rows]" % (limit, total))
    if not lines:
        lines.append("(empty sheet)")
    return "sheet=%s rows=%d\n%s" % (title, total, "\n".join(lines))


@srv.tool("xlsx_set_cell", "Set one cell of an existing workbook, e.g. cell='B2'.",
          {"type": "object", "properties": {"path": {"type": "string"}, "cell": {"type": "string"}, "value": {"type": ["string", "number", "boolean"]}, "sheet": {"type": "string", "default": "Sheet1"}}, "required": ["path", "cell", "value"]})
def xlsx_set_cell(path, cell, value, sheet="Sheet1"):
    p = _in_path(path)
    wb = openpyxl.load_workbook(p)
    if str(sheet) in wb.sheetnames:
        ws = wb[str(sheet)]
    else:
        ws = wb.create_sheet(str(sheet))
    ws[str(cell)] = value
    wb.save(p)
    return "set %s!%s = %r" % (ws.title, cell, value)


@srv.tool("xlsx_sheets", "List sheet names with their dimensions.",
          {"type": "object", "properties": {"path": {"type": "string"}}, "required": ["path"]})
def xlsx_sheets(path):
    p = _in_path(path)
    wb = openpyxl.load_workbook(p, read_only=True)
    try:
        lines = ["%s  rows=%d cols=%d" % (n, wb[n].max_row or 0, wb[n].max_column or 0) for n in wb.sheetnames]
    finally:
        wb.close()
    return "sheets=%d\n%s" % (len(lines), "\n".join(lines) if lines else "(none)")
@srv.tool("pptx_create", "Create a .pptx from slides: [{\"title\": str, \"bullets\": [str]}].",
          {"type": "object", "properties": {"path": {"type": "string"}, "slides": {"type": "array", "items": {"type": "object", "properties": {"title": {"type": "string"}, "bullets": {"type": "array", "items": {"type": "string"}}}}}}, "required": ["path"]})
def pptx_create(path, slides=[]):
    p = _out_path(path)
    prs = Presentation()
    layout = prs.slide_layouts[1]
    count = 0
    for item in slides or []:
        if not isinstance(item, dict):
            raise ValueError("each slide must be an object with title/bullets: %r" % (item,))
        slide = prs.slides.add_slide(layout)
        title = slide.shapes.title
        if title is not None:
            title.text = str(item.get("title", ""))
        frame = slide.placeholders[1].text_frame
        frame.clear()
        for i, bullet in enumerate(item.get("bullets") or []):
            para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            para.text = str(bullet)
            para.level = 0
        count += 1
    prs.save(p)
    return "created %s (slides=%d)" % (p, count)


@srv.tool("pptx_read", "Read slide titles and body text from a .pptx.",
          {"type": "object", "properties": {"path": {"type": "string"}}, "required": ["path"]})
def pptx_read(path):
    p = _in_path(path)
    prs = Presentation(p)
    out = []
    for i, slide in enumerate(prs.slides, 1):
        out.append("[SLIDE %d]" % i)
        title = slide.shapes.title
        title_id = title.shape_id if title is not None else None
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                text = "".join(run.text for run in para.runs).strip() or _cell(para.text).strip()
                if not text:
                    continue
                if title_id is not None and shape.shape_id == title_id:
                    out.append("[TITLE] %s" % text)
                else:
                    out.append("  %s- %s" % ("  " * int(para.level or 0), text))
    return "slides=%d lines=%d\n%s" % (len(prs.slides._sldIdLst), len(out), "\n".join(out) or "(empty)")


@srv.tool("csv_read", "Read a CSV file as 'a | b | c' lines.",
          {"type": "object", "properties": {"path": {"type": "string"}, "max_rows": {"type": "integer", "default": MAX_ROWS}}, "required": ["path"]})
def csv_read(path, max_rows=MAX_ROWS):
    p = _in_path(path)
    limit = max(1, int(max_rows))
    lines = []
    total = 0
    with open(p, "r", encoding="utf-8-sig", errors="replace", newline="") as fh:
        for row in csv.reader(fh):
            total += 1
            if len(lines) < limit:
                lines.append(" | ".join(_cell(v) for v in row))
    if total > limit:
        lines.append("...[truncated at %d of %d rows]" % (limit, total))
    return "rows=%d\n%s" % (total, "\n".join(lines) if lines else "(empty file)")


@srv.tool("csv_write", "Write rows to a CSV file (UTF-8, parents created).",
          {"type": "object", "properties": {"path": {"type": "string"}, "rows": {"type": "array", "items": {"type": "array"}}}, "required": ["path"]})
def csv_write(path, rows=[[]]):
    p = _out_path(path)
    count = 0
    with open(p, "w", encoding="utf-8", newline="") as fh:
        writer = csv.writer(fh)
        for row in rows or []:
            writer.writerow([_cell(v) for v in row])
            count += 1
    return "wrote %s (rows=%d, %d bytes)" % (p, count, os.path.getsize(p))


@srv.tool("office_convert_table_to_xlsx", "Convert a CSV table into an .xlsx sheet.",
          {"type": "object", "properties": {"csv_path": {"type": "string"}, "xlsx_path": {"type": "string"}}, "required": ["csv_path", "xlsx_path"]})
def office_convert_table_to_xlsx(csv_path, xlsx_path):
    src = _in_path(csv_path)
    dst = _out_path(xlsx_path)
    wb = openpyxl.Workbook()
    ws = wb.active
    count = 0
    with open(src, "r", encoding="utf-8-sig", errors="replace", newline="") as fh:
        for row in csv.reader(fh):
            ws.append([_coerce(v) for v in row])
            count += 1
    wb.save(dst)
    return "converted %s -> %s (sheet=%s rows=%d)" % (src, dst, ws.title, count)
_DOCX = os.path.join(SAMPLE_DIR, "sample.docx")
_XLSX = os.path.join(SAMPLE_DIR, "sample.xlsx")
_PPTX = os.path.join(SAMPLE_DIR, "sample.pptx")
_CSV = os.path.join(SAMPLE_DIR, "sample.csv")

# Ordered so the readers run after the writers that produce their inputs.
SAMPLES = {
    "docx_create": {"path": _DOCX, "title": "Sample report", "paragraphs": ["first line", "second line"], "table": [["name", "qty"], ["widget", "3"]]},
    "docx_append": {"path": _DOCX, "paragraphs": ["appended line"]},
    "docx_read": {"path": _DOCX},
    "xlsx_create": {"path": _XLSX, "rows": [["name", "qty"], ["widget", 3]], "sheet": "Sheet1"},
    "xlsx_read": {"path": _XLSX, "sheet": "", "max_rows": 50},
    "xlsx_set_cell": {"path": _XLSX, "cell": "D1", "value": "note", "sheet": "Sheet1"},
    "xlsx_sheets": {"path": _XLSX},
    "pptx_create": {"path": _PPTX, "slides": [{"title": "Intro", "bullets": ["point one", "point two"]}]},
    "pptx_read": {"path": _PPTX},
    "csv_write": {"path": _CSV, "rows": [["name", "qty"], ["widget", 3]]},
    "csv_read": {"path": _CSV, "max_rows": 50},
    "office_convert_table_to_xlsx": {"csv_path": _CSV, "xlsx_path": os.path.join(SAMPLE_DIR, "converted.xlsx")},
}


def build():
    return srv


if __name__ == "__main__":
    srv.run()




